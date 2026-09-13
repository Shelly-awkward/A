# -*- coding: utf-8 -*-
r"""
make_selection_list.py — 事務所檢查「選案用」公司清單產製（自足驗證版 v2）
==========================================================================

用法（先確定 listdata 已跑過該季 run_all.ps1，有 auditor 快照）：
  python make_selection_list.py --year 114 --season 04                 # 全市場
  python make_selection_list.py --year 114 --season 04 --firm 安永     # 只列某事務所
  python make_selection_list.py --year 114 --season 04 -o 清單.xlsx    # 指定輸出檔名

需要套件：openpyxl、pandas、requests、xlrd、msoffcrypto-tool

───────────────────────────────────────────────────────────────────────────
選案邏輯（為什麼這樣設計——給未來的人與 AI）
──────────────────────────────────────────────────────────────────────────
三層漏斗：機器建母體 → 官方分數排序 → 人工圈案。本腳本做前兩層，
第三層（圈案判斷）永遠是檢查員的責任，不寫進程式。

1. 母體 = listdata auditor 快照（MOPS 簽證申報），不是任何「最新狀態」API。
   ○○年度檢查要的是○○Q4 申報當時的簽竢會計師與意見；用最新狀態會把
   已換所客戶漏掉、把新客戶提早算進來，母體就錯了。
2. 分數 = 證交所／櫃買「平時管理權數統計表」（官方參數），不是自訂權重。
   監理選案必須能對外說明依據。財務權數 9 大類＋非財務權數 17 項＝權數總計。
3. 可驗證性是硬要求：輸出檔內嵌兩張原始權數表，清單權數欄掛 INDEX+MATCH
   公式即時取自原始表，「檢核」欄＝總計−原表總計，應全為 0（≠0 自動標紅）。
   任何人拿到檔案都能重驗，不需要相信產製者。
4. 無權數者（公開發行、興櫃、金融、上市 KY）沉底但不刪除——仍是受查客戶，
   只是要用別的依據選案（金融另有監理；上市 KY 需「外國企業版」權數表）。

──────────────────────────────────────────────────────────────────────────
已知陷阱（改動前必讀，都是實測踩過的）
──────────────────────────────────────────────────────────────────────────
★ 權數表欄位陷阱（最重要）：兩份權數表（證交所 xlsx／櫃買 xls）欄 84 的
  「表頭文字」都誤標為「權數總計」，實際資料是**非財務權數合計**；真正的
  權數總計在欄 85（表頭為合併儲存格 CG4:CG5，唯讀模式下讀不到）。
  → 取數一律用固定欄位索引（0-indexed：65=財務合計、83=非財務合計、
    84=權數總計），**不要用表頭文字定位**。此對應已以 208 家安永客戶
    「分項相加＝總計」零誤差驗證（見 權數表\選案權數表分析_交接筆記.md）。
- 權數表版面：第 4–5 列是表頭（第 5 列有分項名稱）、第 6 列起是資料。
  分數項目欄：0-indexed 56–64＝財務 9 類、66–82＝非財務 17 項，
  「主要扣分項目」即由這些欄之得分>0 者組成（名稱截 18 字、分數由高到低）。
- 櫃買 xls 有開檔密碼（114Q4＝3513，寫在檔名），需 msoffcrypto 解密＋
  xlrd encoding_override='cp950' 讀取。證交所版不含 KY（外國企業第一上市）；
  櫃買版含 KY。
- 公司代號型別不一致：auditor 快照是文字（證券期貨商含前導零如 '0009A0'），
  權數表是數字 → join 用 lstrip('0')，Excel 公式用 VALUE() 轉換，
  非數字代號（如 '0009A0'）由 IFERROR 落空白。
- 市場別以權數表歸屬為準（申報檔市場別偶有錯置）。
- 證券期貨商因個體＋合併雙申報，同代號可能出現兩列，屬預期、不去重。
- MOPS ajax_t163sb04/05（EPS／每股淨值）可能被 WAF 擋：已 try/except 降級
  留空；結果會快取到 fin_cache_<季>.json，重跑不重打 MOPS。
- 新承接（初次受任，審計 510 期初餘額）判定＝意見欄含
  「前期由其他會計師查核」，自動寫入備註；權數表本身抓不到此訊號。
- district.csv（承辦人）在 listdata 與 projects\mops-daily 各有一份，
  更新時兩處要同步（本腳本讀 listdata 那份）。
- 輸出檔的公式需 Excel 開啟時重算（已設 fullCalcOnLoad）；
  Google Drive 網頁預覽不會算公式，看到空白屬正常。
- 選案限制（誠實揭露）：權數僅涵蓋公開財報與治理警訊；檢舉、前次檢查缺失、交易所監視名單等監理內部訊號不在模型八，圈案時由檢查員疊加。
"""
import argparse, csv, glob, html, io, json, os, re, sys, time

import requests
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import CellIsRule

LISTDATA = r"C:\Users\dinef\AI\listdata"
HERE = os.path.dirname(os.path.abspath(__file__))
WEIGHTS_DIR = os.path.join(HERE, "權數表")
UA = {"User-Agent": "Mozilla/5.0"}
TYPEKS = ["sii", "otc", "rotc", "pub"]  # 上市/上櫃/興櫃/公開發行
FONT = "Microsoft JhengHei"

# 權數表固定欄位索引（0-indexed；勿改用表頭文字定位，見 docstring 陷阱）
IDX_FIN_ITEMS = range(56, 65)    # 財務 9 類
IDX_FIN_TOTAL = 65               # 財務權數合計
IDX_NF_ITEMS = range(66, 83)     # 非財務 17 項
IDX_NF_TOTAL = 83                # 非財務權數合計（表頭誤標「權數總計」）
IDX_TOTAL = 84                   # 權數總計（表頭在合併儲存格）
HDR_ROWS = 5                     # 表頭列數，資料自第 6 列起


def latest_auditor_file(year, season):
    pat = os.path.join(LISTDATA, f"{year}Q{season}_auditor_*.xlsx")
    files = sorted(glob.glob(pat))
    if not files:
        sys.exit(f"找不到 {pat}，請先到 listdata 執行 run_all.ps1 -Year {year} -Season {season}")
    return files[-1]


def load_auditor(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    return [list(r) for r in rows[1:] if r[1] not in (None, "")]  # 市場別,代號,簡稱,事務所,CPA1,CPA2,日期,意見


def load_basic():
    """公司基本資料：董事長、地址。優先用 listdata 的 company_basic_merged.csv，缺時抓 MOPS opendata。"""
    merged = os.path.join(LISTDATA, "company_basic_merged.csv")
    info = {}
    if os.path.exists(merged):
        with open(merged, encoding="utf-8-sig") as f:
            for row in csv.DictReader(f):
                code = row.get("公司代號", "").strip()
                if code:  # MOPS 來源偶含 HTML 實體字（如 &#20931; = 凃），一律還原
                    info[code] = (html.unescape(row.get("董事長", "")), html.unescape(row.get("住址", "")))
    if info:
        return info
    for sfx in ("L", "O", "R", "P"):
        r = requests.get(f"https://mopsfin.twse.com.tw/opendata/t187ap03_{sfx}.csv", timeout=60, headers=UA)
        r.encoding = "utf-8-sig"
        for row in csv.DictReader(io.StringIO(r.text)):
            code = row.get("公司代號", "").strip()
            if code and code not in info:
                info[code] = (html.unescape(row.get("董事長", "")), html.unescape(row.get("住址", "")))
    return info


def load_officer():
    d = {}
    with open(os.path.join(LISTDATA, "district.csv"), encoding="utf-8-sig") as f:
        for row in csv.DictReader(f):
            d[row["證券代號"].strip()] = row["承辦人員"].strip()
    return d


def fetch_fin(year, season):
    """MOPS 彙總報表：每股盈餘(t163sb04)、每股參考淨值(t163sb05)。結果快取成 json，重跑不重打 MOPS。"""
    cache_path = os.path.join(HERE, f"fin_cache_{year}Q{int(season)}.json")
    if os.path.exists(cache_path):
        print(f"  讀取快取 {os.path.basename(cache_path)}")
        return json.load(io.open(cache_path, encoding="utf-8"))
    fin = {}
    for typek in TYPEKS:
        for form, key, kw in (("ajax_t163sb04", "eps", "每股盈餘"),
                              ("ajax_t163sb05", "nav", "每股參考淨值")):
            try:
                r = requests.post(f"https://mopsov.twse.com.tw/mops/web/{form}",
                                  data={"encodeURIComponent": "1", "step": "1", "firstin": "1", "off": "1",
                                        "isQuery": "Y", "TYPEK": typek, "year": str(year), "season": season},
                                  timeout=90, headers=UA)
                r.encoding = "utf8"
                tables = pd.read_html(io.StringIO(r.text))
            except Exception as e:
                print(f"  ! {typek} {form} 抓取失敗：{e}", file=sys.stderr)
                continue
            for t in tables:
                t.columns = [str(c).replace(" ", "") for c in t.columns]
                if "公司代號" not in t.columns:
                    continue
                hit = [c for c in t.columns if kw in c]
                if not hit:
                    continue
                for _, row in t.iterrows():
                    code = str(row["公司代號"]).split(".")[0]
                    v = row[hit[0]]
                    if pd.notna(v):
                        fin.setdefault(code, {})[key] = v
            print(f"  {typek} {key} 完成")
            time.sleep(2)
    if fin:
        io.open(cache_path, "w", encoding="utf-8").write(json.dumps(fin, ensure_ascii=False))
    return fin


def _breakdown(row_vals, headers5):
    """主要扣分項目(分數)：得分>0 之分項，分數高→低，名稱截 18 字。"""
    items = []
    for c in list(IDX_FIN_ITEMS) + list(IDX_NF_ITEMS):
        v = row_vals[c]
        if isinstance(v, (int, float)) and v > 0:
            name = str(headers5[c] or f"欄{c+1}").strip()
            items.append((v, f"{name[:18]}({int(v)})"))
    items.sort(key=lambda x: -x[0])
    return "；".join(t for _, t in items)


def load_weights(year, season):
    """權數表（證交所+櫃買）。

    回傳：
      w    = {代號: (市場, 財務, 非財務, 總計, 主要扣分項目)}
      raws = {"上市": [[...85欄], ...含表頭5列], "上櫃": [...]}（供內嵌原始表）
    欄位一律用固定索引（見 docstring 陷阱：欄84表頭誤標）。檔案不在就回空。
    ★ 權數表逐年可能改版，固定索引僅對已驗證年度成立——每次載入以
      _verify_layout() 全表驗證「分項加總＝財務合計＋非財務合計＝權數總計」，
      勾稽不過即中止，強迫重新分析新版欄位結構，不得沿用舊索引硬跑。
    """
    def _verify_layout(rows, label):
        bad = total = 0
        for r in rows[HDR_ROWS:]:
            if not isinstance(r[0], (int, float)) or not r[0]:
                continue
            total += 1
            fin_sum = sum(v for c in IDX_FIN_ITEMS if isinstance((v := r[c]), (int, float)))
            nf_sum = sum(v for c in IDX_NF_ITEMS if isinstance((v := r[c]), (int, float)))
            if abs(fin_sum - (r[IDX_FIN_TOTAL] or 0)) > 0.01 or \
               abs(nf_sum - (r[IDX_NF_TOTAL] or 0)) > 0.01 or \
               abs(fin_sum + nf_sum - (r[IDX_TOTAL] or 0)) > 0.01:
                bad += 1
        if total == 0 or bad > 0:
            sys.exit(f"權數表【{label}】欄位勾稽失敗（{bad}/{total} 筆分項加總≠合計）："
                     f"檔案可能改版，請先重新分析欄位結構並更新 IDX_* 常數，勿沿用舊索引。")
        print(f"  權數表【{label}】欄位勾稽通過（{total} 筆分項加總＝合計）")

    w, raws = {}, {}
    q = f"{year}Q{int(season)}"
    twse = glob.glob(os.path.join(WEIGHTS_DIR, f"{q}權數表*上市*.xlsx"))
    if twse:
        wb = openpyxl.load_workbook(twse[0], read_only=True, data_only=True)
        rows = [list(r) for r in wb[wb.sheetnames[0]].iter_rows(values_only=True)]
        wb.close()
        raws["上市"] = rows
        _verify_layout(rows, "上市")
        h5 = rows[HDR_ROWS - 1]
        for r in rows[HDR_ROWS:]:
            if isinstance(r[0], (int, float)):
                w[str(int(r[0]))] = ("上市", r[IDX_FIN_TOTAL] or 0, r[IDX_NF_TOTAL] or 0,
                                     r[IDX_TOTAL] or 0, _breakdown(r, h5))
    otc = glob.glob(os.path.join(WEIGHTS_DIR, f"{q}上櫃*.xls"))
    if otc:
        import msoffcrypto, xlrd
        # 開檔密碼依慣例寫在檔名（如「…加密3513.xls」）；解析不到才退回 114Q4 的 3513
        m = re.search(r"加密(\w+)", os.path.basename(otc[0]))
        pwd = m.group(1) if m else "3513"
        buf = io.BytesIO()
        with open(otc[0], "rb") as f:
            of = msoffcrypto.OfficeFile(f)
            of.load_key(password=pwd)
            of.decrypt(buf)
        buf.seek(0)
        sh = xlrd.open_workbook(file_contents=buf.read(), encoding_override="cp950").sheet_by_index(0)
        rows = [[sh.cell_value(i, c) if sh.cell_value(i, c) != "" else None
                 for c in range(sh.ncols)] for i in range(sh.nrows)]
        raws["上櫃"] = rows
        _verify_layout(rows, "上櫃")
        h5 = rows[HDR_ROWS - 1]
        for r in rows[HDR_ROWS:]:
            if isinstance(r[0], (int, float)) and r[0]:
                w[str(int(r[0]))] = ("上櫃", r[IDX_FIN_TOTAL] or 0, r[IDX_NF_TOTAL] or 0,
                                     r[IDX_TOTAL] or 0, _breakdown(r, h5))
    return w, raws


HDR = ["市場別", "公司代號", "公司簡稱", "承辦人", "董事長", "地址", "事務所", "會計師A", "會計師B",
       "每股盈餘(元)", "每股淨值(元)", "查核意見", "簽證日期", "財務權數", "非財務權數", "權數總計",
       "檢核(與原表總計差)", "主要扣分項目(分數)", "備註"]
WIDTHS = [9, 9, 12, 10, 12, 36, 26, 9, 9, 11, 11, 40, 10, 10, 11, 10, 15, 70, 26]


def _pull_formula(sheet, col_letter, r):
    """INDEX+MATCH 從內嵌原始表取值；代號為文字需 VALUE()；取不到回空白。"""
    return (f"IFERROR(INDEX('{sheet}'!${col_letter}:${col_letter},"
            f"MATCH(VALUE($B{r}),'{sheet}'!$A:$A,0)),\"\")")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--year", required=True, help="民國年，如 114")
    ap.add_argument("--season", required=True, help="01~04")
    ap.add_argument("--firm", default="", help="事務所名稱關鍵字（如：安永），留空=全市場")
    ap.add_argument("-o", "--output", default="", help="輸出檔名（預設自動命名）")
    a = ap.parse_args()
    season = a.season.zfill(2)
    q = f"{a.year}Q{int(season)}"

    src = latest_auditor_file(a.year, season)
    print(f"讀取 auditor 快照：{os.path.basename(src)}")
    rows = load_auditor(src)
    if a.firm:
        rows = [r for r in rows if r[3] and a.firm in str(r[3])]
        print(f"篩選事務所「{a.firm}」：{len(rows)} 筆")
    else:
        print(f"全市場：{len(rows)} 筆")
    if not rows:
        sys.exit("沒有符合的資料。")

    basic = load_basic()
    officer = load_officer()
    print("抓取 MOPS 每股盈餘／每股淨值…")
    fin = fetch_fin(a.year, season)
    weights, raws = load_weights(a.year, season)
    print(f"權數表：{'已併入 ' + str(len(weights)) + ' 家' if weights else '該季檔案不在權數表資料夾，輸出無權數/公式'}")

    data = []
    for m, code, name, firm, cpa1, cpa2, dt, op in (r[:8] for r in rows):
        code = str(code).strip()
        key = code.lstrip("0") or code
        b = basic.get(code) or basic.get(key) or ("", "")
        f = fin.get(key, {})
        wgt = weights.get(key)
        mkt = wgt[0] if wgt else m  # 權數表歸屬優先（修正申報檔偶發的市場別錯置）
        rmk = []
        if op and "前期由其他會計師查核" in str(op):
            rmk.append("新承接（審計510 期初餘額）")
        if "-KY" in str(name):
            rmk.append("KY" if wgt else "上市KY：不在本國權數表，需外國企業版")
        data.append([mkt, code, name, officer.get(code, "") or officer.get(key, ""), b[0], b[1],
                     firm, cpa1, cpa2, f.get("eps"), f.get("nav"), op, dt,
                     int(wgt[1]) if wgt else None, int(wgt[2]) if wgt else None,
                     int(wgt[3]) if wgt else None, None,
                     wgt[4] if wgt else None, "；".join(rmk) or None])
    # 權數總計高→低排序（無分數者沉底），選案直接從頭挑
    data.sort(key=lambda x: (-(x[15] if x[15] is not None else -1), str(x[1])))

    out = openpyxl.Workbook()

    # ── 說明 sheet
    ws0 = out.active
    ws0.title = "說明"
    import datetime
    n_w = sum(1 for d in data if d[15] is not None)
    lines = [
        f"{a.year}Q{int(season)} 選案清單（{a.firm or '全市場'}）— 自足驗證版",
        "",
        f"母體：listdata auditor 快照（MOPS 簽證申報）{'篩「' + a.firm + '」' if a.firm else ''}＝{len(data)} 筆",
        "  （含上市/上櫃/興櫃/公開發行；證券期貨商因個體＋合併雙申報可能重複兩列）。",
        "  會計師A/B、查核意見、簽證日期＝該季申報當時之簽竢資訊（非現時點狀態）。",
        "",
        f"權數分數：證交所／櫃買中心 {q} 平時管理權數統計表（本檔內嵌原始表，逐分可追溯）。",
        f"  匹配 {n_w} 筆有分數；其餘 {len(data) - n_w} 筆（公發、興櫃、金融、上市KY 等）無權數沉底、不刪除。",
        "  財務權數/非財務權數欄＝INDEX+MATCH 公式即時取自內嵌原始表；權數總計＝兩者相加；",
        "  檢核欄＝總計−原表權數總計，應全為 0（≠0 自動標紅）＝開檔即可驗證的勾稽證明。",
        "  ★原始表欄84表頭誤標「權數總計」，實為非財務合計；真總計在欄85（合併表頭）。腳本用固定索引取數。",
        "",
        "主要扣分項目(分數)：由原始表財務 9 類＋非財務 17 項之得分欄彙整（>0 者，分數高→低，名稱截18字）。",
        "備註欄自動標記：新承接（意見含「前期由其他會計師查核」→審計510 期初餘額）；KY。",
        "",
        "其他欄位來源：承辦人＝district.csv（證期局管區）；董事長/地址＝MOPS 公司基本資料；",
        "  EPS/每股淨值＝MOPS 彙總報表 t163sb04/05。",
        "",
        "選案限制（誠實揭露）：權數僆涵蓋公開財報與治理警訊；檢舉、前次檢查缺失、交易所監視名單",
        "  不在模型內。分數是排序輔助，圈案判斷責任在檢查員。",
        f"產製：{datetime.date.today().isoformat()}；腳本＝make_selection_list.py（含邏輯說明，改動前先讀 docstring）。",
    ]
    for i, t in enumerate(lines, 1):
        ws0.cell(i, 1, t).font = Font(name=FONT, size=11, bold=(i == 1))
    ws0.column_dimensions["A"].width = 110

    # ── 清單 sheet
    ws = out.create_sheet(f"{q}選案清單")
    ws.append(HDR)
    for c in range(1, len(HDR) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF", size=10)
        cell.fill = PatternFill("solid", fgColor="305496")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    sheet_sii, sheet_otc = f"權數表原始_上市{q}", f"權數表原始_上櫃{q}"
    body = Font(name=FONT, size=10)
    for d in data:
        ws.append(d)
        r = ws.max_row
        for c in range(1, len(HDR) + 1):
            ws.cell(row=r, column=c).font = body
        ws.cell(row=r, column=10).number_format = "0.00"
        ws.cell(row=r, column=11).number_format = "0.00"
        ws.cell(row=r, column=18).alignment = Alignment(wrap_text=True, vertical="top")
        if d[15] is not None and raws:  # 有權數者掛公式（BN=財務合計 CF=非財務合計 CG=權數總計）
            sheet = sheet_sii if d[0] == "上市" else sheet_otc
            ws.cell(r, 14, "=" + _pull_formula(sheet, "BN", r))
            ws.cell(r, 15, "=" + _pull_formula(sheet, "CF", r))
            ws.cell(r, 16, f"=N{r}+O{r}")
            ws.cell(r, 17, f"=P{r}-" + _pull_formula(sheet, "CG", r))
            for c in (14, 15, 16, 17):
                ws.cell(r, c).font = body

    for i, wd in enumerate(WIDTHS, 1):
        ws.column_dimensions[get_column_letter(i)].width = wd
    ws.freeze_panes = "A2"
    last = len(data) + 1
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HDR))}{last}"
    ws.conditional_formatting.add(
        f"Q2:Q{last}",
        CellIsRule(operator="notEqual", formula=["0"],
                   fill=PatternFill("solid", fgColor="FFC7CE"), font=Font(color="9C0006")))

    # ── 內嵌原始權數表（維持原欄位位置，公式的 BN/CF/CG 才對得上）
    for mkt, sn in (("上市", sheet_sii), ("上櫃", sheet_otc)):
        if mkt not in raws:
            continue
        dst = out.create_sheet(sn)
        for row in raws[mkt]:
            dst.append(row)
        dst.column_dimensions["A"].width = 9
        dst.column_dimensions["B"].width = 14

    out.calculation.fullCalcOnLoad = True
    fn = a.output or f"{q}_選案清單{'_' + a.firm if a.firm else '_全市場'}.xlsx"
    if not os.path.isabs(fn):
        fn = os.path.join(HERE, fn)
    out.save(fn)
    print(f"完成：{fn}（{len(data)} 筆，權數公式 {n_w} 列）")


if __name__ == "__main__":
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    main()
