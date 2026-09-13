# -*- coding: utf-8 -*-
"""
合併總表：安永 114 年度上市櫃財簽客戶 × 兩條確信線

一列一家公司（221 家），並排呈現：
  A 基本資料（財簽事務所／會計師）
  B 永續報告書線（114 年度永續報告書之會計師確信與第三方查證）
  C 路徑圖線（2025 年度溫室氣體排放之確信／查證，出具者由 PDF 逐份判讀）
  D 綜合判斷（重疊態樣）

輸出：安永114客戶_確信總表.xlsx（總表＋統計摘要兩個工作表）、.csv
"""
import sys, os, json, re, time
import pandas as pd
from ghg_verdicts import VERDICTS, FLAGS, FAILED

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
HERE = os.path.dirname(os.path.abspath(__file__))
AUDITOR_XLSX = r"C:\Users\dinef\AI\listdata\114Q04_auditor_20260716.xlsx"
EY_PAT = re.compile(r"安永|Ernst\s*&?\s*Young|EY\s*Taiwan")
TYPE_LABEL = {"cpa": "會計師事務所", "consult": "會計師事務所體系顧問公司",
              "verify": "驗證／查驗機構", "unknown": "未具名"}
NULL_GUID = "00000000-0000-0000-0000-000000000000"


def blank(v):
    return v is None or str(v).strip() in ("", "無", "nan", "None")


def txt(v):
    return "" if blank(v) else str(v).strip()


def load():
    # 財簽名單（安永上市櫃）
    df = pd.read_excel(AUDITOR_XLSX, sheet_name="114Q04")
    df["公司代號"] = df["公司代號"].astype(str).str.strip()
    ey = (df[df["事務所"].astype(str).str.contains("安永")
             & df["市場別"].isin(["上市", "上櫃"])]
          .drop_duplicates("公司代號").sort_values("公司代號"))

    # 永續報告書申報
    mops = {}
    for lst in json.load(open(os.path.join(HERE, "raw_mops_2025.json"), encoding="utf-8")).values():
        for x in lst:
            mops[str(x["code"]).strip()] = x

    # 溫室氣體申報（攤平巢狀 controls）
    ghg = {}
    for lst in json.load(open(os.path.join(HERE, "raw_ghg_2025.json"), encoding="utf-8")).values():
        for x in lst:
            d = {}
            for sec in x.get("sectionDatas") or []:
                for c in sec.get("controls") or []:
                    if c.get("code"):
                        d[c["code"]] = c.get("value")
            ghg[str(x["companyCode"]).strip()] = d
    return ey, mops, ghg


def build():
    ey, mops, ghg = load()
    rows = []
    for _, a in ey.iterrows():
        code = a["公司代號"]
        mk = a["市場別"]
        mt = 0 if mk == "上市" else 1
        r = {
            "公司代號": code,
            "公司名稱": a["公司簡稱"],
            "市場別": mk,
            "財簽事務所": a["事務所"],
            "財簽會計師": "、".join([x for x in [txt(a.get("會計師1")), txt(a.get("會計師2"))] if x]),
        }
        notes = []

        # ── B 永續報告書線 ──────────────────────────────────
        m = mops.get(code)
        if m is None:
            r.update({
                "[永續報告書] 狀態": "尚未出版",
                "[永續報告書] 申報日": "",
                "[永續報告書] 有會計師確信": "不適用",
                "[永續報告書] 確信出具者": "",
                "[永續報告書] 確信準則": "",
                "[永續報告書] 確信等級": "",
                "[永續報告書] 確信範圍": "",
                "[永續報告書] 有第三方查證": "不適用",
                "[永續報告書] 查證機構": "",
                "[永續報告書] 查證標準": "",
                "[永續報告書] 確信者為安永": "不適用",
            })
            sr_overlap = None
        else:
            partner, std = txt(m.get("partnerOrganization")), txt(m.get("stdCompliance"))
            level, scope = txt(m.get("accOptionType")), txt(m.get("aConfirmTarget"))
            vali, tstd, tver = txt(m.get("valiProvider")), txt(m.get("thirdStd")), txt(m.get("thirdPartyVerify"))
            has_a = bool(partner or std or level or scope)
            has_v = bool(vali or tstd or tver)
            is_ey = bool(partner and EY_PAT.search(partner))
            r.update({
                "[永續報告書] 狀態": "已出版",
                "[永續報告書] 申報日": txt(m.get("twCommencementDate")),
                "[永續報告書] 有會計師確信": "是" if has_a else "否",
                "[永續報告書] 確信出具者": partner,
                "[永續報告書] 確信準則": std,
                "[永續報告書] 確信等級": level,
                "[永續報告書] 確信範圍": scope,
                "[永續報告書] 有第三方查證": "是" if has_v else "否",
                "[永續報告書] 查證機構": vali,
                "[永續報告書] 查證標準": tstd,
                "[永續報告書] 確信者為安永": ("是" if is_ey else "否") if has_a else "無確信",
            })
            sr_overlap = is_ey if has_a else False
            if not has_a and not has_v:
                af = txt(m.get("aReportUploadId")) not in ("", NULL_GUID)
                vf = txt(m.get("thirdPartyUploadId")) not in ("", NULL_GUID)
                if not af and not vf:
                    notes.append("永續報告書已出版但確信與查證欄全空、亦未附任何文件，須翻 PDF 附錄複核")

        # ── C 路徑圖線（溫室氣體）──────────────────────────
        g = ghg.get(code)
        if g is None:
            r.update({
                "[溫室氣體] 已申報": "否",
                "[溫室氣體] 範疇一排放量": "", "[溫室氣體] 範疇二排放量": "",
                "[溫室氣體] 範疇一確信範圍": "", "[溫室氣體] 範疇一確信意見": "",
                "[溫室氣體] 有確信報告": "不適用",
                "[溫室氣體] 確信報告出具者": "", "[溫室氣體] 出具者類型": "",
                "[溫室氣體] 確信者為安永": "不適用",
            })
            ghg_overlap = None
        else:
            uid = g.get("greenhouseGasEmissionsAssertionReport")
            has_rep = bool(uid)
            if has_rep:
                v = VERDICTS.get(code)
                if v:
                    issuer, kind, ev = v
                elif code in FAILED:
                    issuer, kind, ev = "（下載失敗，無從辨識）", "unknown", ""
                else:
                    issuer, kind, ev = "（未判讀）", "unknown", ""
                is_ey_g = bool(EY_PAT.search(issuer))
                ghg_overlap = is_ey_g if kind != "unknown" else None
                if code in FLAGS:
                    notes.append("溫室氣體確信報告：" + FLAGS[code])
                if code in FAILED:
                    notes.append("溫室氣體確信報告檔下載失敗，出具者無從辨識")
            else:
                issuer, kind, is_ey_g, ghg_overlap = "", "", False, False
            r.update({
                "[溫室氣體] 已申報": "是",
                "[溫室氣體] 範疇一排放量": txt(g.get("grossScope1GreenhouseGasEmissions")),
                "[溫室氣體] 範疇二排放量": txt(g.get("grossScope2GreenhouseGasEmissions")),
                "[溫室氣體] 範疇一確信範圍": txt(g.get("grossScope1ConfidenceScope")),
                "[溫室氣體] 範疇一確信意見": txt(g.get("grossScope1ConfidenceOpinion")),
                "[溫室氣體] 有確信報告": "是" if has_rep else "否",
                "[溫室氣體] 確信報告出具者": issuer,
                "[溫室氣體] 出具者類型": TYPE_LABEL.get(kind, ""),
                "[溫室氣體] 確信者為安永": ("是" if is_ey_g else "否") if has_rep and kind != "unknown"
                                        else ("無從辨識" if has_rep else "無確信"),
            })

        # ── D 綜合 ────────────────────────────────────────
        if sr_overlap and ghg_overlap:
            state = "兩線皆重疊"
        elif ghg_overlap:
            state = "僅溫室氣體確信重疊"
        elif sr_overlap:
            state = "僅永續報告書確信重疊"
        elif sr_overlap is None and ghg_overlap is None:
            state = "尚無從認定"
        else:
            state = "無重疊"
        r["【綜合】重疊態樣"] = state

        # 還缺什麼才能下最終定論
        gaps = []
        if sr_overlap is None:
            gaps.append("永續報告書未出版")
        if g is None:
            gaps.append("未申報溫室氣體")
        elif not g.get("greenhouseGasEmissionsAssertionReport"):
            gaps.append("溫室氣體未確信")
        elif ghg_overlap is None:
            gaps.append("溫室氣體出具者無從辨識")
        r["【綜合】待補資料"] = "；".join(gaps) if gaps else "兩線齊備"
        r["資料來源"] = (f"永續報告書 https://esggenplus.twse.com.tw/inquiry/report?market={mt}&year=2025&companyCode={code}&lang=zh-TW"
                     f"　｜溫室氣體 https://esggenplus.twse.com.tw/inquiry/info/individual?market={mt}&companyCode={code}&year=2025&lang=zh-TW")
        r["需人工確認"] = ("是：" + "；".join(notes)) if notes else "否"
        rows.append(r)
    return pd.DataFrame(rows)


def stats(df):
    n = len(df)
    out = []
    add = lambda a, b, c="": out.append({"區塊": a, "項目": b, "家數": c})

    def cnt(col, val):
        return int((df[col] == val).sum())

    out.append({"區塊": "母體", "項目": "安永簽證之 114 年度上市櫃客戶", "家數": n})
    out.append({"區塊": "母體", "項目": "　上市", "家數": cnt("市場別", "上市")})
    out.append({"區塊": "母體", "項目": "　上櫃", "家數": cnt("市場別", "上櫃")})

    pub = cnt("[永續報告書] 狀態", "已出版")
    out.append({"區塊": "永續報告書線", "項目": "已出版 114 年度永續報告書", "家數": pub})
    out.append({"區塊": "永續報告書線", "項目": "尚未出版", "家數": n - pub})
    out.append({"區塊": "永續報告書線", "項目": "　其中有會計師確信", "家數": cnt("[永續報告書] 有會計師確信", "是")})
    out.append({"區塊": "永續報告書線", "項目": "　其中有第三方查證", "家數": cnt("[永續報告書] 有第三方查證", "是")})
    out.append({"區塊": "永續報告書線", "項目": "　　確信者為安永（重疊）", "家數": cnt("[永續報告書] 確信者為安永", "是")})

    g = cnt("[溫室氣體] 已申報", "是")
    rep = cnt("[溫室氣體] 有確信報告", "是")
    out.append({"區塊": "路徑圖線（溫室氣體）", "項目": "已申報 2025 年度溫室氣體資料", "家數": g})
    out.append({"區塊": "路徑圖線（溫室氣體）", "項目": "　其中已上傳確信／查證報告", "家數": rep})
    out.append({"區塊": "路徑圖線（溫室氣體）", "項目": "　　出具者為會計師事務所體系", "家數": int(df["[溫室氣體] 出具者類型"].isin(["會計師事務所", "會計師事務所體系顧問公司"]).sum())})
    out.append({"區塊": "路徑圖線（溫室氣體）", "項目": "　　出具者為驗證／查驗機構", "家數": cnt("[溫室氣體] 出具者類型", "驗證／查驗機構")})
    out.append({"區塊": "路徑圖線（溫室氣體）", "項目": "　　出具者未具名／無從辨識", "家數": cnt("[溫室氣體] 出具者類型", "未具名")})
    out.append({"區塊": "路徑圖線（溫室氣體）", "項目": "　　　確信者為安永（重疊）", "家數": cnt("[溫室氣體] 確信者為安永", "是")})

    for k, v in df["【綜合】重疊態樣"].value_counts().items():
        out.append({"區塊": "綜合－重疊態樣", "項目": k, "家數": int(v)})
    for k, v in df["【綜合】待補資料"].value_counts().items():
        out.append({"區塊": "綜合－待補資料", "項目": k, "家數": int(v)})
    out.append({"區塊": "綜合", "項目": "需人工確認", "家數": int(df["需人工確認"].str.startswith("是").sum())})

    # 出具者排行
    for col, blk in [("[永續報告書] 查證機構", "永續報告書－查證機構"),
                     ("[永續報告書] 確信出具者", "永續報告書－確信出具者"),
                     ("[溫室氣體] 確信報告出具者", "溫室氣體－確信報告出具者")]:
        vc = df[df[col].astype(str).str.strip() != ""][col].value_counts()
        for k, v in vc.items():
            out.append({"區塊": blk, "項目": k, "家數": int(v)})
    return pd.DataFrame(out)


if __name__ == "__main__":
    df = build()
    st = stats(df)
    xlsx = os.path.join(HERE, "安永114客戶_確信總表.xlsx")
    with pd.ExcelWriter(xlsx, engine="openpyxl") as w:
        df.to_excel(w, index=False, sheet_name="總表")
        st.to_excel(w, index=False, sheet_name="統計摘要")
        ws = w.sheets["總表"]
        ws.freeze_panes = "C2"
        ws.auto_filter.ref = ws.dimensions
        from openpyxl.utils import get_column_letter
        widths = [10, 14, 8, 22, 16,
                  14, 12, 16, 26, 22, 10, 34, 16, 26, 24, 16,
                  12, 16, 16, 22, 26, 14, 40, 22, 16,
                  22, 26, 60, 60]
        for i, wd in enumerate(widths[:len(df.columns)], 1):
            ws.column_dimensions[get_column_letter(i)].width = wd
        from openpyxl.styles import Alignment, Font, PatternFill
        fills = {"[永續報告書]": "FFF2E5", "[溫室氣體]": "E5F1FF", "【綜合】": "E8F5E9"}
        for i, c in enumerate(df.columns, 1):
            cell = ws.cell(row=1, column=i)
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            cell.font = Font(bold=True)
            for k, col in fills.items():
                if c.startswith(k):
                    cell.fill = PatternFill("solid", fgColor=col)
        ws2 = w.sheets["統計摘要"]
        for col, wd in [("A", 24), ("B", 52), ("C", 10)]:
            ws2.column_dimensions[col].width = wd
        ws2.freeze_panes = "A2"
    df.to_csv(os.path.join(HERE, "安永114客戶_確信總表.csv"), index=False, encoding="utf-8-sig")
    print(f"→ {xlsx}（{len(df)} 列 × {len(df.columns)} 欄）")
    print(st.to_string(index=False))
