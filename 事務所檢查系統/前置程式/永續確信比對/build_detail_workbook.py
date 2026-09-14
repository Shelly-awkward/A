# -*- coding: utf-8 -*-
"""
明細活頁簿：把兩條線的完整明細合成一個 xlsx 的兩張 sheet。

  sheet「永續報告書明細」 221 列 — 安永全部上市櫃客戶 × 114 年度永續報告書確信／查證
  sheet「溫室氣體明細」    79 列 — 其中已上傳 2025 年度溫室氣體確信報告者 × PDF 逐份判讀出具者

與 安永114客戶_確信總表.xlsx 的分工：
  總表   一列一家、兩線並排，用來看全貌與統計
  明細表 保留各線完整欄位（確信範圍原文、佐證摘要、判讀依據），用來查個案為什麼這樣認定

用法：python build_detail_workbook.py
"""
import sys, os
import pandas as pd
from openpyxl.utils import get_column_letter
from openpyxl.styles import Alignment, Font

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
HERE = os.path.dirname(os.path.abspath(__file__))

import build_ey_assurance as sr
import build_ghg_overlap as gh

SHEETS = [
    ("永續報告書明細", [10, 14, 8, 22, 16, 14, 30, 16, 28, 24, 10, 40, 16, 28, 26, 16, 22, 40, 60, 50]),
    ("溫室氣體明細", [10, 14, 8, 22, 16, 44, 24, 14, 16, 26, 18, 18, 18, 60, 60, 60]),
]


def main():
    df_sr, year = sr.build(2025, refresh=False)
    df_gh = gh.build()

    out = os.path.join(HERE, "安永114客戶_確信明細表.xlsx")
    with pd.ExcelWriter(out, engine="openpyxl") as w:
        for (name, widths), df in zip(SHEETS, [df_sr, df_gh]):
            df.to_excel(w, index=False, sheet_name=name)
            ws = w.sheets[name]
            ws.freeze_panes = "C2"
            ws.auto_filter.ref = ws.dimensions
            for i, wd in enumerate(widths[:len(df.columns)], 1):
                ws.column_dimensions[get_column_letter(i)].width = wd
            for i in range(1, len(df.columns) + 1):
                c = ws.cell(row=1, column=i)
                c.font = Font(bold=True)
                c.alignment = Alignment(wrap_text=True, vertical="center")
    print(f"→ {out}")
    print(f"   永續報告書明細 {len(df_sr)} 列 × {len(df_sr.columns)} 欄")
    print(f"   溫室氣體明細   {len(df_gh)} 列 × {len(df_gh.columns)} 欄")


if __name__ == "__main__":
    main()
